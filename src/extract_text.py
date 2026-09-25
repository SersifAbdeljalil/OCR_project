"""
extract_text.py - Extraction du texte selon le format (SANS OCR a cette etape).

    PDF          : PyMuPDF, page par page ; une page sans couche texte est
                   marquee "OCR requis" (seuil commun : SEUIL_CARACTERES_PAGE).
    Images       : aucune extraction, toutes les pages marquees "OCR requis".
    DOCX / DOTX  : python-docx (paragraphes, tableaux, zones de texte,
                   en-tetes, pieds de page) ; secours par lecture du XML.
    XLS / XLSX   : xlrd / openpyxl, toutes les feuilles, cellules non vides
                   ligne par ligne, dates Excel converties en AAAA-MM-JJ.

REGLE : extraire() ne leve JAMAIS d'exception. Un fichier protege, corrompu ou
d'un format inconnu donne illisible=True et une alerte "illisible : <raison>".
Les alertes ne contiennent jamais de texte du document ni de message systeme
(qui peut contenir un chemin) : seulement le type de probleme.
"""

# --- Imports ---------------------------------------------------------------
import io
import re
import zipfile
from dataclasses import dataclass, field
from datetime import date, datetime, time
from pathlib import Path
from xml.etree import ElementTree

import pymupdf

from src.config import SEUIL_CARACTERES_PAGE

# --- Constantes ------------------------------------------------------------
EXT_PDF = {".pdf"}
EXT_IMAGES = {".jpg", ".jpeg", ".png", ".tif", ".tiff", ".bmp", ".webp"}
EXT_WORD = {".docx", ".dotx"}
EXT_XLS = {".xls"}
EXT_XLSX = {".xlsx"}

STATUT_TEXTE = "texte"
STATUT_OCR = "OCR requis"
SEPARATEUR_PAGES = "\n\n"
SEPARATEUR_CELLULES = " | "

# Types de contenu Word : un modele (.dotx) est refuse tel quel par python-docx
TYPE_MODELE = "application/vnd.openxmlformats-officedocument.wordprocessingml.template.main+xml"
TYPE_DOCUMENT = "application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"
NS_W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


@dataclass
class ResultatExtraction:
    """Resultat commun a tous les formats."""
    texte: str = ""
    nb_pages: int = 0                  # pages (PDF, image) ou feuilles (Excel), 1 pour Word
    pages_ocr: list = field(default_factory=list)       # numeros de pages (a partir de 1)
    statut_pages: list = field(default_factory=list)    # "texte" / "OCR requis" par page
    # Texte natif page par page (PDF, images : "" pour une page a OCR) ; vide pour
    # Word / Excel (pas de pages). Sert a fusionner le texte OCR dans l'ordre.
    textes_pages: list = field(default_factory=list)
    methode: str = "aucune"
    illisible: bool = False
    alertes: list = field(default_factory=list)


def _illisible(resultat: ResultatExtraction, raison: str) -> ResultatExtraction:
    resultat.illisible = True
    resultat.alertes.append(f"illisible : {raison}")
    return resultat


# --- 1. PDF ----------------------------------------------------------------
def _nb_caracteres_visibles(texte: str) -> int:
    return sum(1 for c in texte if not c.isspace())


def extraire_pdf(chemin: Path) -> ResultatExtraction:
    """Texte des pages natives ; les autres sont listees dans pages_ocr."""
    r = ResultatExtraction(methode="pymupdf")
    try:
        with pymupdf.open(chemin) as doc:
            if doc.needs_pass:
                return _illisible(r, "pdf protege par mot de passe")
            r.nb_pages = doc.page_count
            if r.nb_pages == 0:
                return _illisible(r, "pdf sans aucune page")
            textes = []
            for numero, page in enumerate(doc, start=1):
                texte = page.get_text()
                if _nb_caracteres_visibles(texte) >= SEUIL_CARACTERES_PAGE:
                    r.statut_pages.append(STATUT_TEXTE)
                    textes.append(texte.strip())
                    r.textes_pages.append(texte.strip())
                else:
                    r.statut_pages.append(STATUT_OCR)
                    r.pages_ocr.append(numero)
                    r.textes_pages.append("")
    except Exception as err:
        return _illisible(r, f"fichier corrompu ou invalide ({type(err).__name__})")
    r.texte = SEPARATEUR_PAGES.join(textes)
    return r


# --- 2. Images -------------------------------------------------------------
def extraire_image(chemin: Path) -> ResultatExtraction:
    """Aucune extraction : on verifie seulement que l'image s'ouvre et on compte
    ses pages (un TIFF peut en avoir plusieurs). Tout passera en OCR."""
    r = ResultatExtraction(methode="image")
    try:
        with pymupdf.open(chemin) as doc:
            r.nb_pages = doc.page_count
            # PyMuPDF "ouvre" meme un faux fichier image : seul le decodage revele
            # une image corrompue. Rendu minuscule (10 %) de chaque page : peu couteux.
            for page in doc:
                page.get_pixmap(matrix=pymupdf.Matrix(0.1, 0.1))
    except Exception as err:
        return _illisible(r, f"image corrompue ou invalide ({type(err).__name__})")
    if r.nb_pages == 0:
        return _illisible(r, "image vide")
    r.pages_ocr = list(range(1, r.nb_pages + 1))
    r.statut_pages = [STATUT_OCR] * r.nb_pages
    r.textes_pages = [""] * r.nb_pages
    return r


# --- 3. Word (DOCX / DOTX) -------------------------------------------------
def _texte_tableau(tableau) -> list:
    """Une ligne de texte par ligne du tableau ; cellules fusionnees non repetees."""
    lignes = []
    for rangee in tableau.rows:
        cellules, vues = [], set()
        for cellule in rangee.cells:
            if id(cellule._tc) in vues:           # cellule fusionnee deja lue
                continue
            vues.add(id(cellule._tc))
            texte = cellule.text.strip()
            if texte:
                cellules.append(texte)
        if cellules:
            lignes.append(SEPARATEUR_CELLULES.join(cellules))
    return lignes


def _texte_bloc(conteneur) -> list:
    """Paragraphes et tableaux d'un conteneur (corps, en-tete...), dans l'ordre."""
    from docx.table import Table
    from docx.text.paragraph import Paragraph
    lignes = []
    for enfant in conteneur._element.iterchildren():
        if enfant.tag == f"{NS_W}p":
            texte = Paragraph(enfant, conteneur).text.strip()
            if texte:
                lignes.append(texte)
        elif enfant.tag == f"{NS_W}tbl":
            lignes += _texte_tableau(Table(enfant, conteneur))
    return lignes


def _texte_zones_de_texte(element) -> list:
    """Zones de texte (w:txbxContent), ignorees par python-docx. Word les enregistre
    souvent deux fois (version moderne + version de secours) : doublons retires."""
    lignes = []
    for zone in element.iter(f"{NS_W}txbxContent"):
        for p in zone.iter(f"{NS_W}p"):
            texte = "".join(t.text or "" for t in p.iter(f"{NS_W}t")).strip()
            if texte:
                lignes.append(texte)
    return list(dict.fromkeys(lignes))


def _texte_docx(document) -> str:
    """En-tetes, corps (+ zones de texte), pieds de page. Un en-tete repete dans
    plusieurs sections n'est ecrit qu'une fois."""
    entetes, pieds = [], []
    for section in document.sections:
        for partie in (section.first_page_header, section.header, section.even_page_header):
            if not partie.is_linked_to_previous:
                entetes += _texte_bloc(partie) + _texte_zones_de_texte(partie._element)
        for partie in (section.first_page_footer, section.footer, section.even_page_footer):
            if not partie.is_linked_to_previous:
                pieds += _texte_bloc(partie) + _texte_zones_de_texte(partie._element)
    corps = _texte_bloc(document._body) + _texte_zones_de_texte(document.element.body)
    lignes = list(dict.fromkeys(entetes)) + corps + list(dict.fromkeys(pieds))
    return "\n".join(lignes)


def _docx_depuis_modele(chemin: Path):
    """Copie EN MEMOIRE d'un .dotx dont le type de contenu est change en
    'document' : python-docx accepte alors de l'ouvrir. Le fichier n'est pas modifie."""
    tampon = io.BytesIO()
    with zipfile.ZipFile(chemin) as source, zipfile.ZipFile(tampon, "w") as copie:
        for element in source.infolist():
            donnees = source.read(element.filename)
            if element.filename == "[Content_Types].xml":
                donnees = donnees.replace(TYPE_MODELE.encode(), TYPE_DOCUMENT.encode())
            copie.writestr(element, donnees)
    tampon.seek(0)
    return tampon


def texte_xml_word(chemin: Path) -> str:
    """Secours : lit directement le XML (en-tetes, document, pieds de page).
    Un paragraphe (w:p) = une ligne. Moins precis (tableaux a plat)."""
    lignes = []
    with zipfile.ZipFile(chemin) as z:
        noms = z.namelist()
        parties = (sorted(n for n in noms if re.fullmatch(r"word/header\d*\.xml", n))
                   + ["word/document.xml"]
                   + sorted(n for n in noms if re.fullmatch(r"word/footer\d*\.xml", n)))
        for nom in parties:
            if nom not in noms:
                continue
            racine = ElementTree.fromstring(z.read(nom))
            for p in racine.iter(f"{NS_W}p"):
                texte = "".join(t.text or "" for t in p.iter(f"{NS_W}t")).strip()
                if texte and (not lignes or lignes[-1] != texte):
                    lignes.append(texte)
    return "\n".join(lignes)


def extraire_word(chemin: Path) -> ResultatExtraction:
    """python-docx ; pour un modele .dotx refuse, copie en memoire avec le type
    'document' ; en dernier recours, lecture brute du XML."""
    import docx
    r = ResultatExtraction(nb_pages=1)
    try:
        r.texte, r.methode = _texte_docx(docx.Document(str(chemin))), "python-docx"
        return r
    except ValueError:
        pass                                      # format modele (.dotx) refuse
    except Exception as err:
        r.alertes.append(f"python-docx a echoue ({type(err).__name__}) : lecture du XML")
    try:
        r.texte = _texte_docx(docx.Document(_docx_depuis_modele(chemin)))
        r.methode = "python-docx (modele)"
        return r
    except Exception:
        pass
    try:
        r.texte, r.methode = texte_xml_word(chemin), "xml"
        return r
    except Exception as err:
        return _illisible(r, f"fichier Word corrompu ou protege ({type(err).__name__})")


# --- 4. Excel (XLS / XLSX) -------------------------------------------------
def _texte_cellule(valeur) -> str:
    """Date -> AAAA-MM-JJ (avec l'heure si elle n'est pas minuit) ;
    1200.0 -> 1200 ; texte nettoye."""
    if isinstance(valeur, datetime):
        return (valeur.date().isoformat() if valeur.time() == time(0, 0)
                else valeur.isoformat(sep=" ", timespec="minutes"))
    if isinstance(valeur, (date, time)):
        return valeur.isoformat()
    if isinstance(valeur, float) and valeur.is_integer():
        return str(int(valeur))
    if isinstance(valeur, bool):
        return "VRAI" if valeur else "FAUX"
    return str(valeur).strip()


def _texte_feuilles(feuilles) -> str:
    """feuilles : liste de (nom, lignes de valeurs). Une ligne de texte par ligne
    non vide, precedee du nom de la feuille."""
    blocs = []
    for nom, lignes in feuilles:
        texte_lignes = []
        for valeurs in lignes:
            cellules = [_texte_cellule(v) for v in valeurs if v is not None]
            cellules = [c for c in cellules if c]
            if cellules:
                texte_lignes.append(SEPARATEUR_CELLULES.join(cellules))
        if texte_lignes:
            blocs.append(f"[Feuille : {nom}]\n" + "\n".join(texte_lignes))
    return SEPARATEUR_PAGES.join(blocs)


def extraire_xlsx(chemin: Path) -> ResultatExtraction:
    import openpyxl
    r = ResultatExtraction(methode="openpyxl")
    try:
        # data_only=True : valeur calculee des formules (et non la formule)
        classeur = openpyxl.load_workbook(chemin, read_only=True, data_only=True)
        try:
            feuilles = [(f.title, list(f.iter_rows(values_only=True)))
                        for f in classeur.worksheets]
        finally:
            classeur.close()
    except Exception as err:
        return _illisible(r, f"classeur corrompu ou protege ({type(err).__name__})")
    r.nb_pages, r.texte = len(feuilles), _texte_feuilles(feuilles)
    return r


def extraire_xls(chemin: Path) -> ResultatExtraction:
    import xlrd
    r = ResultatExtraction(methode="xlrd")
    try:
        classeur = xlrd.open_workbook(str(chemin))
        feuilles = []
        for f in classeur.sheets():
            lignes = []
            for i in range(f.nrows):
                valeurs = []
                for cellule in f.row(i):
                    if cellule.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
                        continue
                    if cellule.ctype == xlrd.XL_CELL_DATE:     # date Excel = nombre de jours
                        valeurs.append(xlrd.xldate_as_datetime(cellule.value,
                                                               classeur.datemode))
                    elif cellule.ctype == xlrd.XL_CELL_BOOLEAN:
                        valeurs.append(bool(cellule.value))
                    elif cellule.ctype == xlrd.XL_CELL_ERROR:
                        continue
                    else:
                        valeurs.append(cellule.value)
                lignes.append(valeurs)
            feuilles.append((f.name, lignes))
    except Exception as err:
        return _illisible(r, f"classeur corrompu ou protege ({type(err).__name__})")
    r.nb_pages, r.texte = len(feuilles), _texte_feuilles(feuilles)
    return r


# --- 5. Point d'entree -----------------------------------------------------
def extraire(chemin) -> ResultatExtraction:
    """Extrait le texte d'un fichier. Ne leve JAMAIS d'exception."""
    try:
        chemin = Path(chemin)
        ext = chemin.suffix.lower()
        if not chemin.is_file():
            return _illisible(ResultatExtraction(), "fichier introuvable")
        if ext in EXT_PDF:
            r = extraire_pdf(chemin)
        elif ext in EXT_IMAGES:
            r = extraire_image(chemin)
        elif ext in EXT_WORD:
            r = extraire_word(chemin)
        elif ext in EXT_XLSX:
            r = extraire_xlsx(chemin)
        elif ext in EXT_XLS:
            r = extraire_xls(chemin)
        else:
            return _illisible(ResultatExtraction(), f"format inconnu ({ext or 'sans extension'})")
    except Exception as err:                       # filet de securite final
        return _illisible(ResultatExtraction(), f"erreur inattendue ({type(err).__name__})")

    # Document lisible mais vide (sans page a passer en OCR) : on le signale
    if not r.illisible and not r.texte.strip() and not r.pages_ocr:
        r.alertes.append("aucun texte extrait")
    return r
