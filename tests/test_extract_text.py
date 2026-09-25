"""
test_extract_text.py - Tests de l'extraction de texte (src/extract_text.py).

Tous les fichiers sont INVENTES et generes dans tmp_path : aucun document reel.
Lancement (venv active, depuis OCR_PROJECT) :  python -m pytest -v
"""

import zipfile
from datetime import date, datetime

import docx
import openpyxl
import pymupdf
import pytest

from src.extract_text import (STATUT_OCR, STATUT_TEXTE, TYPE_DOCUMENT, TYPE_MODELE,
                              extraire, texte_xml_word)

LIGNE = "Facture inventée FA-2026-00042 : montant HT 1 000,00 DH, TVA 200,00 DH."


# --- Fabrication de fichiers inventes --------------------------------------
def creer_pdf(chemin, pages, **options_sauvegarde):
    """pages : liste de textes ; "" = page sans couche texte (comme un scan)."""
    doc = pymupdf.open()
    for texte in pages:
        page = doc.new_page()
        if texte:
            page.insert_text((72, 72), texte, fontsize=9)
    doc.save(chemin, **options_sauvegarde)
    doc.close()
    return chemin


def creer_image(chemin):
    pix = pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, 20, 20), False)
    pix.clear_with(255)
    pix.save(chemin)
    return chemin


def creer_docx(chemin):
    d = docx.Document()
    d.sections[0].header.paragraphs[0].text = "En-tête Société Exemple"
    d.sections[0].footer.paragraphs[0].text = "Pied de page : ICE inventé"
    d.add_paragraph("Contrat de prestation inventé")
    tableau = d.add_table(rows=2, cols=3)
    for i, valeurs in enumerate([("Désignation", "Qté", "Prix"),
                                 ("Maintenance", "2", "500,00")]):
        for j, v in enumerate(valeurs):
            tableau.cell(i, j).text = v
    d.add_paragraph("Fait à Casablanca")
    d.save(chemin)
    return chemin


def en_modele_dotx(source, cible):
    """Transforme un .docx en .dotx (type de contenu 'template')."""
    with zipfile.ZipFile(source) as zs, zipfile.ZipFile(cible, "w") as zc:
        for e in zs.infolist():
            donnees = zs.read(e.filename)
            if e.filename == "[Content_Types].xml":
                donnees = donnees.replace(TYPE_DOCUMENT.encode(), TYPE_MODELE.encode())
            zc.writestr(e, donnees)
    return cible


def creer_xlsx(chemin):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Factures"
    ws.append(["Numéro", "Date", "Montant"])
    ws.append(["FA-001", datetime(2026, 9, 15), 1200.0])
    ws.append([None, None, None])                     # ligne vide : ignoree
    ws.append(["FA-002", date(2026, 10, 1), 99.5])
    ws2 = wb.create_sheet("Notes")
    ws2["B3"] = "Remarque inventée"
    wb.save(chemin)
    return chemin


# --- 1. PDF ----------------------------------------------------------------
def test_pdf_natif(tmp_path):
    r = extraire(creer_pdf(tmp_path / "natif.pdf", [LIGNE]))
    assert r.methode == "pymupdf" and not r.illisible and r.alertes == []
    assert r.nb_pages == 1 and r.pages_ocr == [] and r.statut_pages == [STATUT_TEXTE]
    assert "FA-2026-00042" in r.texte and "inventée" in r.texte


def test_pdf_sans_texte(tmp_path):
    r = extraire(creer_pdf(tmp_path / "scan.pdf", ["", ""]))
    assert r.pages_ocr == [1, 2] and r.statut_pages == [STATUT_OCR, STATUT_OCR]
    assert r.texte == "" and not r.illisible
    assert r.alertes == []                   # pas "aucun texte" : l'OCR s'en chargera


def test_pdf_mixte(tmp_path):
    r = extraire(creer_pdf(tmp_path / "mixte.pdf", [LIGNE, "", "Page trois : " + LIGNE]))
    assert r.nb_pages == 3 and r.pages_ocr == [2]
    assert r.statut_pages == [STATUT_TEXTE, STATUT_OCR, STATUT_TEXTE]
    assert "Page trois" in r.texte


def test_pdf_page_sous_le_seuil(tmp_path):
    """Un numero de page seul (< 50 caracteres) ne compte pas comme du texte."""
    r = extraire(creer_pdf(tmp_path / "presque_vide.pdf", ["Page 1"]))
    assert r.pages_ocr == [1]


def test_pdf_protege(tmp_path):
    chemin = creer_pdf(tmp_path / "protege.pdf", [LIGNE],
                       encryption=pymupdf.PDF_ENCRYPT_AES_256,
                       owner_pw="proprio", user_pw="secret")
    r = extraire(chemin)
    assert r.illisible and r.alertes == ["illisible : pdf protege par mot de passe"]
    assert r.texte == ""


def test_pdf_corrompu(tmp_path):
    chemin = tmp_path / "corrompu.pdf"
    chemin.write_bytes(b"ceci n'est pas un pdf")
    r = extraire(chemin)
    assert r.illisible and r.alertes[0].startswith("illisible : fichier corrompu")


# --- 2. Images -------------------------------------------------------------
@pytest.mark.parametrize("nom", ["photo.png", "photo.JPG"])
def test_image_ocr_requis(tmp_path, nom):
    chemin = creer_image(tmp_path / nom.lower())
    chemin = chemin.rename(tmp_path / nom)
    r = extraire(chemin)
    assert r.methode == "image" and r.texte == "" and not r.illisible
    assert r.nb_pages == 1 and r.pages_ocr == [1] and r.statut_pages == [STATUT_OCR]


def test_image_corrompue(tmp_path):
    chemin = tmp_path / "photo.jpg"
    chemin.write_bytes(b"pas une image")
    r = extraire(chemin)
    assert r.illisible and "image corrompue" in r.alertes[0]


# --- 3. Word ---------------------------------------------------------------
def test_docx_paragraphes_tableau_entete_pied(tmp_path):
    r = extraire(creer_docx(tmp_path / "contrat.docx"))
    assert r.methode == "python-docx" and not r.illisible
    lignes = r.texte.split("\n")
    assert lignes[0] == "En-tête Société Exemple"                # en-tete d'abord
    assert lignes[-1] == "Pied de page : ICE inventé"            # pied de page a la fin
    # Ordre du corps respecte : paragraphe, tableau (ligne par ligne), paragraphe
    i = lignes.index("Contrat de prestation inventé")
    assert lignes[i + 1:i + 4] == ["Désignation | Qté | Prix",
                                   "Maintenance | 2 | 500,00", "Fait à Casablanca"]


def test_dotx_modele(tmp_path):
    source = creer_docx(tmp_path / "source.docx")
    chemin = en_modele_dotx(source, tmp_path / "modele.dotx")
    with pytest.raises(ValueError):
        docx.Document(str(chemin))                  # python-docx refuse le modele...
    r = extraire(chemin)                            # ... mais l'extraction passe
    assert r.methode == "python-docx (modele)" and not r.illisible
    assert "Maintenance | 2 | 500,00" in r.texte


def test_lecture_xml_de_secours(tmp_path):
    texte = texte_xml_word(creer_docx(tmp_path / "c.docx"))
    for attendu in ("En-tête Société Exemple", "Contrat de prestation inventé",
                    "Maintenance", "Pied de page : ICE inventé"):
        assert attendu in texte


def test_docx_corrompu(tmp_path):
    chemin = tmp_path / "abime.docx"
    chemin.write_bytes(b"pas un zip")
    r = extraire(chemin)
    assert r.illisible and r.alertes[-1].startswith("illisible : fichier Word corrompu")


def test_docx_vide(tmp_path):
    chemin = tmp_path / "vide.docx"
    docx.Document().save(chemin)
    r = extraire(chemin)
    assert not r.illisible and r.texte == "" and r.alertes == ["aucun texte extrait"]


# --- 4. Excel --------------------------------------------------------------
def test_xlsx_feuilles_lignes_dates(tmp_path):
    r = extraire(creer_xlsx(tmp_path / "factures.xlsx"))
    assert r.methode == "openpyxl" and r.nb_pages == 2 and not r.illisible
    assert r.texte == ("[Feuille : Factures]\n"
                       "Numéro | Date | Montant\n"
                       "FA-001 | 2026-09-15 | 1200\n"
                       "FA-002 | 2026-10-01 | 99.5\n\n"
                       "[Feuille : Notes]\n"
                       "Remarque inventée")


@pytest.mark.parametrize("nom", ["abime.xlsx", "abime.xls"])
def test_excel_corrompu(tmp_path, nom):
    chemin = tmp_path / nom
    chemin.write_bytes(b"pas un classeur")
    r = extraire(chemin)
    assert r.illisible and r.alertes[0].startswith("illisible : classeur corrompu")


# --- 5. Cas generaux -------------------------------------------------------
@pytest.mark.parametrize("nom, raison", [
    ("notes.txt", "format inconnu (.txt)"),
    ("ancien.doc", "format inconnu (.doc)"),
    ("sans_extension", "format inconnu (sans extension)"),
])
def test_format_inconnu(tmp_path, nom, raison):
    chemin = tmp_path / nom
    chemin.write_bytes(b"contenu")
    r = extraire(chemin)
    assert r.illisible and r.alertes == [f"illisible : {raison}"]


def test_fichier_introuvable(tmp_path):
    r = extraire(tmp_path / "absent.pdf")
    assert r.illisible and r.alertes == ["illisible : fichier introuvable"]


def test_jamais_d_exception(tmp_path):
    """Des fichiers vides de chaque format : aucune exception, jamais."""
    for nom in ("a.pdf", "b.png", "c.docx", "d.dotx", "e.xls", "f.xlsx", "g.xyz"):
        chemin = tmp_path / nom
        chemin.write_bytes(b"")
        r = extraire(chemin)
        assert r.illisible


def test_alertes_sans_contenu_ni_chemin(tmp_path):
    """Les alertes ne recopient ni le texte ni le chemin (qui peut contenir un nom)."""
    dossier = tmp_path / "Prenom_Nom"
    dossier.mkdir()
    chemin = dossier / "Prenom_Nom.pdf"
    chemin.write_bytes(b"%PDF-1.4 abime Prenom_Nom")
    r = extraire(chemin)
    assert r.illisible
    assert not any("Prenom" in a or str(tmp_path) in a for a in r.alertes)
